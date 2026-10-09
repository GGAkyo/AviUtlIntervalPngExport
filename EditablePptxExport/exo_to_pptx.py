"""AviUtl 1 EXO -> editable PowerPoint prototype. Never modifies input files."""
from __future__ import annotations
import argparse
import configparser
import ctypes
from ctypes import wintypes
from dataclasses import dataclass, field
from functools import lru_cache
from io import BytesIO
import json
import math
import os
from pathlib import Path
import re
import sys

class ConversionError(Exception):
    pass

@dataclass
class Object:
    id: int
    start: int
    end: int
    layer: int
    metadata: dict
    filters: list[dict]

@dataclass
class Project:
    path: Path
    width: int
    height: int
    length: int
    rate: int
    scale: int
    objects: list[Object]

@dataclass
class Report:
    source: str
    output: str
    frames: list[int]
    warnings: list[dict] = field(default_factory=list)
    slides: list[dict] = field(default_factory=list)
    status: str = "checking"
    notes: list[str] = field(default_factory=lambda: [
        "試作版です。AviUtlとPowerPointの文字描画差による字幅・行間・位置の差は残ります。",
        "EXO内の1始まりのフレーム番号を使用します。AUP全体の時刻と一致するとは限りません。",
        "背景色は指定値です。EXOに含まれない本体側のフィルタ・背景・レイヤー状態は復元できません。"
    ])

    def warn(self, obj, message):
        item = {"object": obj.id if obj else None, "message": message}
        if item not in self.warnings:
            self.warnings.append(item)

    def save(self, path):
        from dataclasses import asdict
        with Path(path).open("x", encoding="utf-8") as stream:
            json.dump(asdict(self), stream, ensure_ascii=False, indent=2)

def integer(value, label):
    try:
        return int(value)
    except (ValueError, TypeError) as exc:
        raise ConversionError(f"{label}: 整数が必要です ({value!r})") from exc

def read_project(path):
    path = Path(path).resolve()
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp932")
    cfg = configparser.ConfigParser(interpolation=None, strict=True)
    cfg.optionxform = str
    try:
        cfg.read_string(text)
    except configparser.Error as exc:
        raise ConversionError(f"EXOの解析に失敗しました: {exc}") from exc
    if "exedit" not in cfg:
        raise ConversionError("[exedit] がありません。AviUtl 1のEXOを指定してください。")
    h = cfg["exedit"]
    values = {k: integer(h.get(k), k) for k in ("width", "height", "length", "rate", "scale")}
    if any(v <= 0 for v in values.values()):
        raise ConversionError("解像度・長さ・フレームレートは正数が必要です。")
    objects = []
    for name in cfg.sections():
        if name == "exedit" or re.fullmatch(r"\d+\.\d+", name):
            continue
        if not name.isdecimal():
            raise ConversionError(f"未対応のセクション: [{name}]")
        d = dict(cfg[name])
        start, end, layer = [integer(d.get(k), f"[{name}] {k}") for k in ("start", "end", "layer")]
        if start < 1 or end < start or layer < 1:
            raise ConversionError(f"[{name}] の表示期間またはレイヤーが不正です。")
        parts = sorted((s for s in cfg.sections() if re.fullmatch(re.escape(name) + r"\.\d+", s)),
                       key=lambda s: int(s.split(".")[1]))
        objects.append(Object(int(name), start, end, layer, d, [dict(cfg[s]) for s in parts]))
    return Project(path, **values, objects=objects)

def decode_text(value):
    try:
        result = bytes.fromhex(value).decode("utf-16-le").split("\0", 1)[0]
    except (ValueError, UnicodeError) as exc:
        raise ConversionError("テキストのUTF-16LEデータが不正です。") from exc
    result = result.replace("\r\n", "\n").replace("\r", "\n")
    if any(ord(c) < 32 and c not in "\n\t" for c in result):
        raise ConversionError("テキストに未対応の制御文字があります。")
    return result

def number(d, key, default=0):
    value = d.get(key, str(default))
    if "," in value:
        raise ConversionError(f"{key} の移動設定は未対応です ({value})。")
    try:
        result = float(value)
    except ValueError as exc:
        raise ConversionError(f"{key} の数値が不正です ({value})。") from exc
    if not math.isfinite(result):
        raise ConversionError(f"{key} に有限の数値が必要です。")
    return result

def color(value):
    if not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        raise ConversionError(f"色は6桁のRGB値で指定してください: {value}")
    return value.upper()

def sample_frames(project, interval=30, start=1, end=None, frames=None):
    if frames is not None:
        chosen = sorted(set(frames))
        if not chosen or any(f < 1 or f > project.length for f in chosen):
            raise ConversionError("フレーム番号は1からEXOのlengthの範囲で指定してください。")
    else:
        end = project.length if end is None else end
        if interval < 1 or start < 1 or end < start or end > project.length:
            raise ConversionError("間隔は1以上、開始・終了はEXOの範囲内で指定してください。")
        chosen = list(range(start, end + 1, interval))
    if len(chosen) > 2000:
        raise ConversionError("一度に作成できるスライドは2000枚までです。間隔または範囲を変更してください。")
    return chosen

def image_path(value, project, asset_dir, source_base=None):
    value = value.strip().strip('"')
    if not value:
        raise ConversionError("画像ファイルのパスがありません。")
    p = Path(value)
    base = Path(source_base) if source_base else project.path.parent
    candidates = [Path(asset_dir) / p.name] if asset_dir else [p if p.is_absolute() else base / p]
    for candidate in candidates:
        if candidate.is_file():
            return candidate.resolve()
    raise ConversionError(f"画像が見つかりません: {value}（画像フォルダで移動先を指定できます）")

def spatial_effects(obj, src, draw, report):
    """Only source-space effects before standard drawing can be reproduced."""
    effects = []
    for index, f in enumerate(obj.filters):
        if f is src or f is draw:
            continue
        if (f.get('_name') in ('クリッピング', '反転')
                and obj.filters.index(src) < index < obj.filters.index(draw)
                and src['_name'] in ('画像ファイル', '図形')):
            effects.append(f)
        else:
            report.warn(obj, f"効果「{f.get('_name', '不明')}」は未対応です。近似変換では効果を省略します。")
    return effects

def checkbox(d, key):
    value = number(d, key)
    if value not in (0, 1):
        raise ConversionError(f'{key} は0または1が必要です。')
    return bool(value)

def apply_spatial_effects(item, effects, obj, report):
    # Track the surviving rectangle in the original bitmap, separately from
    # its visible center relative to AviUtl's rotation origin. This also keeps
    # clip -> flip and flip -> clip distinct without discarding source pixels.
    width, height = item['width']/item['scale'], item['height']/item['scale']
    original_width, original_height = width, height
    left = top = right = bottom = 0
    dx = dy = 0
    flip_h = flip_v = False
    for effect in effects:
        if effect['_name'] == 'クリッピング':
            values = [number(effect, k) for k in ('左', '上', '右', '下')]
            if any(v < 0 or not v.is_integer() for v in values):
                raise ConversionError('クリッピング量は0以上の整数が必要です。')
            l, t, r, b = values
            if l+r >= width or t+b >= height:
                raise ConversionError('クリッピング後の幅・高さは1ピクセル以上必要です。')
            recenter = checkbox(effect, '中心の位置を変更')
            if not recenter:
                dx += (l-r)/2
                dy += (t-b)/2
            left += r if flip_h else l
            right += l if flip_h else r
            top += b if flip_v else t
            bottom += t if flip_v else b
            width -= l+r
            height -= t+b
            known = {'_name', '左', '上', '右', '下', '中心の位置を変更'}
        else:
            h, v = checkbox(effect, '左右反転'), checkbox(effect, '上下反転')
            if h:
                flip_h = not flip_h
                dx = -dx
            if v:
                flip_v = not flip_v
                dy = -dy
            for key in ('輝度反転', '色相反転', '透明度反転'):
                if checkbox(effect, key):
                    report.warn(obj, f'反転の「{key}」は未対応です。近似変換では省略します。')
            known = {'_name', '左右反転', '上下反転', '輝度反転', '色相反転', '透明度反転'}
        for key in effect:
            if key not in known:
                report.warn(obj, f"{effect['_name']}の項目 {key} は未対応です。")
    item.update(width=width*item['scale'], height=height*item['scale'],
                offset_x=dx*item['scale'], offset_y=dy*item['scale'],
                flip_h=flip_h, flip_v=flip_v,
                crop=(left/original_width, top/original_height,
                      right/original_width, bottom/original_height))
    return item

def prepare(obj, project, report, asset_dir=None, source_base=None):
    if any(obj.metadata.get(k, "0") != "0" for k in ("camera", "clipping")):
        raise ConversionError("カメラ制御・クリッピングは未対応です。")
    if "chain" in obj.metadata or "group" in obj.metadata:
        raise ConversionError("中間点・グループを含むオブジェクトは未対応です。")
    if obj.metadata.get("overlay", "1") != "1":
        raise ConversionError("overlay=0 のオブジェクトは未対応です。")
    sources = [f for f in obj.filters if f.get("_name") in ("テキスト", "画像ファイル", "図形")]
    draws = [f for f in obj.filters if f.get("_name") == "標準描画"]
    if len(sources) != 1 or len(draws) != 1:
        raise ConversionError("テキスト／画像ファイル／図形と標準描画の組み合わせ以外は未対応です。")
    src, draw = sources[0], draws[0]
    effects = spatial_effects(obj, src, draw, report)
    if number(draw, "Z") != 0 or number(draw, "blend") != 0:
        raise ConversionError("Z座標・通常以外の合成モードは未対応です。")
    for key in draw:
        if key not in {"_name", "X", "Y", "Z", "拡大率", "透明度", "回転", "blend"}:
            report.warn(obj, f"標準描画の項目 {key} は未対応です。")
    scale = number(draw, "拡大率", 100) / 100
    opacity = 1 - number(draw, "透明度") / 100
    if scale <= 0 or not 0 <= opacity <= 1:
        raise ConversionError("拡大率は正数、透明度は0～100が必要です。")
    item = dict(id=obj.id, layer=obj.layer, start=obj.start, end=obj.end,
                x=project.width/2 + number(draw, "X"), y=project.height/2 + number(draw, "Y"),
                scale=scale, opacity=opacity, rotation=number(draw, "回転") % 360)
    if src["_name"] == "図形":
        figure_type = integer(src.get("type", "0"), "図形 type")
        # Built-in EXO IDs: 0 = background, 2 = rectangle.
        # A nonempty name can refer to a custom figure; never guess its geometry.
        if src.get("name", "").strip():
            raise ConversionError("名前付き・カスタム図形は未対応です。")
        if figure_type not in (0, 2):
            raise ConversionError(f"図形 type={figure_type} は未対応です（背景・塗りつぶし四角形に対応）。")
        size = number(src, "サイズ", 100)
        aspect = number(src, "縦横比")
        line = number(src, "ライン幅", 4000)
        if size <= 0 or not -100 < aspect < 100 or line < 0:
            raise ConversionError("図形のサイズ・縦横比・ライン幅が対応範囲外です。")
        if figure_type == 0:
            # Background fills the source canvas; its size/aspect/line controls
            # do not determine the canvas dimensions.
            width, height = project.width, project.height
        else:
            if line < size / 2:
                raise ConversionError("枠線のみの四角形は未対応です。塗りつぶし四角形を使用してください。")
            width = size * (1 - max(aspect, 0) / 100)
            height = size * (1 + min(aspect, 0) / 100)
        for key in src:
            if key not in {"_name", "サイズ", "縦横比", "ライン幅", "type", "color", "name"}:
                report.warn(obj, f"図形の項目 {key} は未対応です。")
        item.update(kind="shape", figure_type=figure_type,
                    width=width*scale, height=height*scale, anchor_x=.5, anchor_y=.5,
                    color=color(src.get("color", "ffffff")))
        return apply_spatial_effects(item, effects, obj, report)
    if src["_name"] == "画像ファイル":
        from PIL import Image
        path = image_path(src.get("file", ""), project, asset_dir, source_base)
        with Image.open(path) as im:
            if getattr(im, "n_frames", 1) != 1:
                raise ConversionError("アニメーション画像は未対応です。")
            width, height = im.size
            buffer = BytesIO()
            rgba = im.convert("RGBA")
            if opacity != 1:
                rgba.putalpha(rgba.getchannel("A").point(lambda a: round(a * opacity)))
            rgba.save(buffer, format="PNG")
        for key in src:
            if key not in {"_name", "file"}:
                report.warn(obj, f"画像ファイルの項目 {key} は未対応です。")
        item.update(kind="image", width=width*scale, height=height*scale,
                    data=buffer.getvalue(), path=str(path), anchor_x=.5, anchor_y=.5)
        return apply_spatial_effects(item, effects, obj, report)
    for key in ("表示速度", "文字毎に個別オブジェクト", "移動座標上に表示する", "自動スクロール"):
        if number(src, key) != 0:
            raise ConversionError(f"テキストの {key} は未対応です。")
    text = decode_text(src.get("text", ""))
    if re.search(r"<[^>]*>", text) or "\t" in text:
        raise ConversionError("テキスト制御タグ・タブは未対応です。")
    if number(src, "type") != 0:
        report.warn(obj, "影付き・縁取り文字は近似変換で通常の文字に置き換えます。")
    for key in ("autoadjust", "monospace"):
        if number(src, key) != 0:
            report.warn(obj, f"テキストの {key} は未対応です。")
    known = {"_name", "サイズ", "表示速度", "文字毎に個別オブジェクト", "移動座標上に表示する",
             "自動スクロール", "B", "I", "type", "autoadjust", "soft", "monospace", "align",
             "spacing_x", "spacing_y", "precision", "color", "color2", "font", "text"}
    for key in src:
        if key not in known:
            report.warn(obj, f"テキストの項目 {key} は未対応です。")
    align = integer(src.get("align", "0"), "align")
    if not 0 <= align <= 8:
        raise ConversionError("縦書きなどの文字揃えは未対応です。")
    size = number(src, "サイズ", 24)
    if not 1 <= size <= 2000:
        raise ConversionError("文字サイズは1～2000に対応しています。")
    font = src.get("font", "MS UI Gothic")
    bold, italic = bool(number(src, "B")), bool(number(src, "I"))
    lines = text.split("\n")
    widths, line_height, actual_font = measure_text(tuple(lines), font, round(size), bold, italic)
    if actual_font.casefold() != font.casefold():
        report.warn(obj, f"フォント「{font}」を計測できず「{actual_font}」で代用しました。表示が異なります。")
    spacing_x, spacing_y = number(src, "spacing_x"), number(src, "spacing_y")
    width = max((w + max(0, len(line)-1)*spacing_x for w, line in zip(widths, lines)), default=1)
    line_step = line_height + spacing_y
    if line_step <= 0 or width < 0:
        raise ConversionError("負の文字幅・行送りには対応していません。")
    height = line_height + max(0, len(lines)-1)*line_step
    item.update(kind="text", text=text, font=font, bold=bold, italic=italic, size=size*scale,
                spacing_x=spacing_x*scale, line_step=line_step*scale,
                color=color(src.get("color", "ffffff")), width=max(1, width)*scale,
                height=max(1, height)*scale, anchor_x=(align % 3)/2, anchor_y=(align // 3)/2,
                align=align % 3)
    return item

@lru_cache(maxsize=1024)
def measure_text(lines, font, size, bold=False, italic=False):
    """Measure Windows fonts in pixels without rasterizing the slide."""
    if os.name != "nt":
        raise ConversionError("この試作版の文字計測にはWindowsが必要です。")
    gdi = ctypes.WinDLL("gdi32", use_last_error=True)
    gdi.CreateCompatibleDC.argtypes = [wintypes.HDC]
    gdi.CreateCompatibleDC.restype = wintypes.HDC
    gdi.CreateFontW.argtypes = [ctypes.c_int]*5 + [wintypes.DWORD]*8 + [wintypes.LPCWSTR]
    gdi.CreateFontW.restype = wintypes.HANDLE
    gdi.SelectObject.argtypes = [wintypes.HDC, wintypes.HANDLE]
    gdi.SelectObject.restype = wintypes.HANDLE
    gdi.DeleteObject.argtypes = [wintypes.HANDLE]
    gdi.DeleteDC.argtypes = [wintypes.HDC]
    class SIZE(ctypes.Structure):
        _fields_ = [("cx", wintypes.LONG), ("cy", wintypes.LONG)]
    gdi.GetTextExtentPoint32W.argtypes = [wintypes.HDC, wintypes.LPCWSTR, ctypes.c_int, ctypes.POINTER(SIZE)]
    gdi.GetTextFaceW.argtypes = [wintypes.HDC, ctypes.c_int, wintypes.LPWSTR]
    dc = gdi.CreateCompatibleDC(None)
    handle = gdi.CreateFontW(-size, 0, 0, 0, 700 if bold else 400, int(italic), 0, 0, 1, 0, 0, 4, 0, font)
    if not dc or not handle:
        if dc: gdi.DeleteDC(dc)
        if handle: gdi.DeleteObject(handle)
        raise ConversionError("Windowsでのフォント計測に失敗しました。")
    old = gdi.SelectObject(dc, handle)
    try:
        widths = []
        height = size
        for line in lines:
            extent = SIZE()
            value = line or " "
            if not gdi.GetTextExtentPoint32W(dc, value, len(value.encode("utf-16-le"))//2, ctypes.byref(extent)):
                raise ConversionError("文字サイズを計測できませんでした。")
            widths.append(float(extent.cx) if line else 0.)
            height = max(height, extent.cy)
        face = ctypes.create_unicode_buffer(128)
        gdi.GetTextFaceW(dc, len(face), face)
        # GetTextFaceW may return a localized name for the requested family.
        if face.value.casefold() != font.casefold():
            class LOGFONT(ctypes.Structure):
                _fields_ = [(k, wintypes.LONG) for k in
                            ("height", "width", "escapement", "orientation", "weight")] + [
                            (k, wintypes.BYTE) for k in
                            ("italic", "underline", "strikeout", "charset", "out", "clip", "quality", "pitch")] + [
                            ("face", wintypes.WCHAR * 32)]
            lf = LOGFONT()
            lf.charset = 1
            lf.face = font[:31]
            found = []
            CALLBACK = ctypes.WINFUNCTYPE(ctypes.c_int, ctypes.c_void_p, ctypes.c_void_p,
                                         wintypes.DWORD, wintypes.LPARAM)
            def visit(*_):
                found.append(True)
                return 0
            callback = CALLBACK(visit)
            gdi.EnumFontFamiliesExW.argtypes = [wintypes.HDC, ctypes.POINTER(LOGFONT),
                                               CALLBACK, wintypes.LPARAM, wintypes.DWORD]
            gdi.EnumFontFamiliesExW(dc, ctypes.byref(lf), callback, 0, 0)
            if found:
                return widths, float(height), font
        return widths, float(height), face.value
    finally:
        gdi.SelectObject(dc, old)
        gdi.DeleteObject(handle)
        gdi.DeleteDC(dc)

def top_left(item):
    # PPT rotates around shape center; EXO text may be anchored at an edge.
    dx = (.5-item["anchor_x"])*item["width"] + item.get('offset_x', 0)
    dy = (.5-item["anchor_y"])*item["height"] + item.get('offset_y', 0)
    a = math.radians(item["rotation"])
    cx = item["x"] + dx*math.cos(a) - dy*math.sin(a)
    cy = item["y"] + dx*math.sin(a) + dy*math.cos(a)
    return cx-item["width"]/2, cy-item["height"]/2

def build_presentation(project, items, report, slide_width, background):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width = Inches(slide_width)
    prs.slide_height = round(prs.slide_width * project.height/project.width)
    unit = prs.slide_width / project.width
    emu = lambda pixels: round(pixels * unit)
    pt = lambda pixels: pixels * unit / 12700
    for frame in report.frames:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
        entry = {"frame": frame, "seconds": (frame-1)*project.scale/project.rate, "objects": []}
        for item in items:
            if not item["start"] <= frame <= item["end"] or item["opacity"] == 0:
                continue
            left, top = top_left(item)
            bounds = [emu(left), emu(top), max(1,emu(item["width"])), max(1,emu(item["height"]))]
            if item["kind"] == "image":
                shape = slide.shapes.add_picture(BytesIO(item["data"]), *bounds)
                (shape.crop_left, shape.crop_top,
                 shape.crop_right, shape.crop_bottom) = item['crop']
                # DrawingML flips preserve the editable crop and full bitmap.
                xfrm = shape._element.spPr.get_or_add_xfrm()
                if item['flip_h']:
                    xfrm.set('flipH', '1')
                if item['flip_v']:
                    xfrm.set('flipV', '1')
            elif item["kind"] == "shape":
                shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *bounds)
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(item["color"])
                shape.line.fill.background()
                # Remove the theme effect reference (which can add a shadow).
                style = shape._element.find(
                    "{http://schemas.openxmlformats.org/presentationml/2006/main}style")
                if style is not None:
                    shape._element.remove(style)
                if item["opacity"] != 1:
                    alpha = OxmlElement("a:alpha")
                    alpha.set("val", str(round(item["opacity"]*100000)))
                    shape._element.spPr.find(
                        "{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")[0].append(alpha)
            else:
                shape = slide.shapes.add_textbox(*bounds)
                tf = shape.text_frame
                tf.clear()
                tf.word_wrap = False
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.vertical_anchor = MSO_ANCHOR.TOP
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                for i, line in enumerate(item["text"].split("\n")):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = (PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT)[item["align"]]
                    p.space_before = p.space_after = Pt(0)
                    p.line_spacing = Pt(pt(item["line_step"]))
                    run = p.add_run()
                    run.text = line
                    run.font.name = item["font"]
                    run.font.size = Pt(pt(item["size"]))
                    run.font.bold, run.font.italic = item["bold"], item["italic"]
                    run.font.color.rgb = RGBColor.from_string(item["color"])
                    rpr = run._r.get_or_add_rPr()
                    rpr.set("lang", "ja-JP")
                    rpr.set("spc", str(round(pt(item["spacing_x"])*100)))
                    rpr.set("kern", "400000")
                    ea = OxmlElement("a:ea")
                    ea.set("typeface", item["font"])
                    rpr.append(ea)
                    if item["opacity"] != 1:
                        alpha = OxmlElement("a:alpha")
                        alpha.set("val", str(round(item["opacity"]*100000)))
                        rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")[0].append(alpha)
            shape.name = f"EXO {item['id']} / Layer {item['layer']} / {item['kind']}"
            shape.rotation = item["rotation"]
            entry["objects"].append({"id": item["id"], "kind": item["kind"],
                                     "left_px": left, "top_px": top, "width_px": item["width"],
                                     "height_px": item["height"], "rotation": item["rotation"],
                                     "crop": list(item.get('crop', (0,0,0,0))),
                                     "flip_h": item.get('flip_h', False),
                                     "flip_v": item.get('flip_v', False)})
        slide.notes_slide.notes_text_frame.text = f"EXO frame: {frame}\nTime: {entry['seconds']:.6f} sec\nSource: {project.path.name}"
        report.slides.append(entry)
    return prs

def resolve_project_citations(project, references=None, source_base=None, default_references=None, citation_state=None, citation_scene=0):
    from citations import resolve, has_commands, load_entries, CitationError
    fields = [(obj, f) for obj in project.objects for f in obj.filters if f.get('_name') == 'テキスト']
    rows = [(index, obj.start, obj.layer, decode_text(f.get('text', ''))) for index, (obj, f) in enumerate(fields)]
    try:
        from citation_state import state_path, load_state, recover
        candidate = Path(default_references) if default_references else (Path(source_base) if source_base else project.path.parent) / 'references.json'
        metadata = Path(citation_state) if citation_state else state_path(references or candidate)
        templates, embedded, warnings = recover(rows, load_state(metadata), citation_scene)
        rows = [(identity, start, layer, templates[identity]) for identity, start, layer, text in rows]
        if not any(has_commands(text) for _, _, _, text in rows): return {}
        if warnings:
            raise CitationError('直接編集された引用文章があります。引用管理画面でキーを確認してください。')
        entries = load_entries(references) if references else None
        defaults = load_entries(candidate) if candidate.is_file() and not references else None
        defaults = {**(defaults or {}), **embedded}
        rendered, numbers, _, _ = resolve(rows, entries, defaults)
        for index, (obj, f) in enumerate(fields):
            f['text'] = (rendered[index] + '\0').encode('utf-16-le').hex()
        return numbers
    except (CitationError, OSError, ValueError) as exc:
        raise ConversionError(f'引用番号を更新できません: {exc}') from exc


def convert(source, output, *, interval=30, start=1, end=None, frames=None,
            asset_dir=None, allow_approx=False, slide_width=13.333333, background="000000",
            source_base=None, references=None, default_references=None, citation_state=None, citation_scene=0):
    output = Path(output).resolve()
    report_path = output.with_suffix(".report.json")
    if output.suffix.lower() != ".pptx":
        raise ConversionError("出力ファイルの拡張子は .pptx にしてください。")
    if output.exists() or report_path.exists():
        raise ConversionError("出力PPTXまたはレポートが既に存在します。別の出力名を指定してください。")
    if not math.isfinite(slide_width) or not 1 <= slide_width <= 56:
        raise ConversionError("スライド幅は1～56インチで指定してください。")
    background = color(background)
    project = read_project(source)
    citation_numbers = resolve_project_citations(project, references, source_base, default_references, citation_state, citation_scene)
    if not 1 <= slide_width*project.height/project.width <= 56:
        raise ConversionError("スライドの縦横比がPowerPointのサイズ制限を超えています。")
    selected = sample_frames(project, interval, start, end, frames)
    report = Report(str(project.path), str(output), selected)
    if citation_numbers:
        report.notes.append("引用番号: " + json.dumps(citation_numbers, ensure_ascii=False))
    items = []
    for obj in sorted(project.objects, key=lambda o: (o.layer, o.id)):
        if not any(obj.start <= f <= obj.end for f in selected):
            continue
        try:
            items.append(prepare(obj, project, report, asset_dir, source_base))
        except (ConversionError, OSError, ValueError) as exc:
            report.warn(obj, f"オブジェクトを省略: {exc}")
    if report.warnings and not allow_approx:
        report.status = "blocked"
        report.save(report_path)
        raise ConversionError(f"未対応の設定が見つかったためPPTXは作成していません。\n詳細: {report_path}\n近似変換を許可する場合は別の出力名で --allow-approx を指定してください。")
    prs = build_presentation(project, items, report, slide_width, background)
    buffer = BytesIO()
    prs.save(buffer)
    with output.open("xb") as stream:
        stream.write(buffer.getvalue())
    report.status = "approximate" if report.warnings else "converted"
    report.save(report_path)
    return report

def main():
    parser = argparse.ArgumentParser(description="AviUtl 1 EXOを編集可能なPPTXに変換します（試作版）")
    parser.add_argument("source", type=Path, help="拡張編集から出力したEXO")
    parser.add_argument("-o", "--output", type=Path, required=True)
    parser.add_argument("--interval", type=int, default=30, help="フレーム間隔 (既定: 30)")
    parser.add_argument("--start", type=int, default=1, help="EXOの開始フレーム (1始まり)")
    parser.add_argument("--end", type=int)
    parser.add_argument("--frames", help="明示的なフレーム番号。例: 1,31,61（範囲・間隔指定より優先）")
    parser.add_argument("--citation-state", type=Path, help="外部保存した引用管理ファイル")
    parser.add_argument("--citation-scene", type=int, default=0)
    parser.add_argument("--references", type=Path, help="文献キーと書誌情報のJSON")
    parser.add_argument("--asset-dir", type=Path, help="画像移動先のフォルダ（同じファイル名で解決）")
    parser.add_argument("--allow-approx", action="store_true", help="未対応効果の省略・未対応オブジェクトの除外を許可")
    parser.add_argument("--slide-width", type=float, default=13.333333, help="スライド幅 (インチ)")
    parser.add_argument("--background", default="000000", help="背景のRGB値 (既定: 000000)")
    args = parser.parse_args()
    try:
        frames = [integer(v.strip(), "--frames") for v in args.frames.split(",")] if args.frames else None
        report = convert(args.source, args.output, interval=args.interval, start=args.start, end=args.end,
                         frames=frames, asset_dir=args.asset_dir, allow_approx=args.allow_approx,
                         slide_width=args.slide_width, background=args.background, references=args.references, citation_state=args.citation_state, citation_scene=args.citation_scene)
    except (ConversionError, OSError, UnicodeError, ImportError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"作成: {args.output} / {len(report.frames)}枚 / 警告: {len(report.warnings)}件")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
