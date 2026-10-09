"""Regression checks for editable content, geometry, timing and safe failure."""
from pathlib import Path
import json
import tempfile
import unittest
from io import BytesIO
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from exo_to_pptx import (convert, decode_text, read_project, sample_frames,
                          top_left, ConversionError)

HEADER = "[exedit]\nwidth=1280\nheight=720\nrate=30\nscale=1\nlength=90\n"

def text_object(text, ident=0, start=1, end=90, layer=2, x=-500, y=-240, align=0, extra=""):
    return f"""
[{ident}]
start={start}
end={end}
layer={layer}
overlay=1
camera=0
[{ident}.0]
_name=テキスト
サイズ=48
表示速度=0.0
文字毎に個別オブジェクト=0
移動座標上に表示する=0
自動スクロール=0
B=0
I=0
type=0
autoadjust=0
soft=1
monospace=0
align={align}
spacing_x=0
spacing_y=0
precision=1
color=ffffff
color2=000000
font=MS Gothic
text={text.encode('utf-16-le').hex()}0000
[{ident}.1]
_name=標準描画
X={x}
Y={y}
Z=0.0
拡大率=100.00
透明度=0.0
回転=0.00
blend=0
{extra}
"""

def image_object(path="asset.png", ident=1, layer=1):
    return f"""
[{ident}]
start=1
end=90
layer={layer}
overlay=1
camera=0
[{ident}.0]
_name=画像ファイル
file={path}
[{ident}.1]
_name=標準描画
X=200
Y=50
Z=0
拡大率=200
透明度=25
回転=15
blend=0
"""


def figure_object(kind=2, size=200, aspect=-50, line=4000, ident=0):
    return f"""
[{ident}]
start=1
end=90
layer=1
overlay=1
camera=0
[{ident}.0]
_name=図形
サイズ={size}
縦横比={aspect}
ライン幅={line}
type={kind}
color=b0b0ff
name=
[{ident}.1]
_name=標準描画
X=100
Y=-100
Z=0
拡大率=150
透明度=25
回転=30
blend=0
"""

class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=Path(__file__).parent)
        self.folder = Path(self.tmp.name)
        self.src = self.folder / "日本語 sample.exo"
        self.out = self.folder / "result.pptx"
        Image.new("RGBA", (100,60), (50,140,200,200)).save(self.folder/"asset.png")

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, body, encoding="cp932"):
        self.src.write_text(HEADER+body, encoding=encoding)
        return self.src

    def test_unicode_and_padding(self):
        text = "日本語 😀\r\nABC"
        self.assertEqual(decode_text((text+"\0\0").encode("utf-16-le").hex()), "日本語 😀\nABC")
        with self.assertRaises(ConversionError):
            decode_text("xyz")

    def test_encoding_and_boundaries(self):
        self.write(text_object("テスト"))
        p = read_project(self.src)
        self.assertEqual(sample_frames(p), [1,31,61])
        self.assertEqual(sample_frames(p, interval=30, start=30, end=90), [30,60,90])
        self.assertEqual(sample_frames(p, frames=[90,1,1]), [1,90])
        for interval in (0,-1):
            with self.assertRaises(ConversionError):
                sample_frames(p, interval=interval)
        with self.assertRaises(ConversionError):
            sample_frames(p, frames=[0])

    def test_editable_japanese_timing_layer_geometry(self):
        self.write(text_object("コピー可能な文字\n2行目", end=30)
                   +image_object()+text_object("次のページ", ident=2, start=31, end=90))
        report = convert(self.src, self.out)
        self.assertEqual(report.warnings, [])
        prs = Presentation(self.out)
        self.assertEqual(len(prs.slides), 3)
        first = prs.slides[0]
        self.assertEqual(first.shapes[0].shape_type, MSO_SHAPE_TYPE.PICTURE)
        self.assertEqual(first.shapes[1].text, "コピー可能な文字\n2行目")
        self.assertEqual(prs.slides[1].shapes[1].text, "次のページ")
        unit = prs.slide_width / 1280
        self.assertAlmostEqual(first.shapes[1].left/unit, 140, places=3)
        self.assertAlmostEqual(first.shapes[1].top/unit, 120, places=3)
        self.assertAlmostEqual(first.shapes[0].width/unit, 200, places=3)
        self.assertAlmostEqual(first.shapes[0].height/unit, 120, places=3)
        self.assertEqual(first.shapes[0].rotation, 15)
        with Image.open(BytesIO(first.shapes[0].image.blob)) as im:
            self.assertEqual(im.getpixel((0,0))[3], 150)
        self.assertIn('typeface="MS Gothic"', first.shapes[1]._element.xml)
        self.assertIn("EXO frame: 31", prs.slides[1].notes_slide.notes_text_frame.text)

    def test_missing_asset_strict_and_approximate(self):
        self.write(image_object("missing.png"))
        with self.assertRaises(ConversionError):
            convert(self.src, self.out)
        self.assertFalse(self.out.exists())
        self.assertEqual(json.loads(self.out.with_suffix(".report.json").read_text("utf8"))["status"], "blocked")
        other = self.folder/"approx.pptx"
        report = convert(self.src, other, allow_approx=True)
        self.assertEqual(report.status, "approximate")
        self.assertEqual(len(Presentation(other).slides[0].shapes), 0)

    def test_existing_files_not_modified(self):
        self.write(text_object("保存"))
        original = self.src.read_bytes()
        self.out.write_bytes(b"must not change")
        with self.assertRaises(ConversionError):
            convert(self.src, self.out)
        self.assertEqual(self.out.read_bytes(), b"must not change")
        self.assertEqual(self.src.read_bytes(), original)

    def test_existing_report_not_overwritten(self):
        self.write(text_object("保存"))
        self.out.with_suffix(".report.json").write_text("keep")
        with self.assertRaises(ConversionError):
            convert(self.src, self.out)
        self.assertFalse(self.out.exists())

    def test_unsupported_motion_is_not_silently_frozen(self):
        self.write(text_object("移動", x="0,100,1"))
        with self.assertRaises(ConversionError):
            convert(self.src, self.out)
        log = self.out.with_suffix(".report.json").read_text("utf8")
        self.assertIn("移動設定", log)

    def test_effect_approximation_is_explicit(self):
        effect = "[0.2]\n_name=ぼかし\n範囲=10\n"
        self.write(text_object("効果")+effect)
        report = convert(self.src, self.out, allow_approx=True)
        self.assertTrue(any("ぼかし" in w["message"] for w in report.warnings))
        self.assertEqual(Presentation(self.out).slides[0].shapes[0].text, "効果")

    def test_unselected_unsupported_object_does_not_block(self):
        self.write(text_object("表示", end=30)+text_object("移動", ident=2, start=31, x="0,100,1"))
        self.assertFalse(convert(self.src, self.out, frames=[1]).warnings)

    def test_anchor_rotation(self):
        item = dict(x=100, y=100, width=80, height=20, anchor_x=0, anchor_y=0, rotation=90)
        left, top = top_left(item)
        self.assertAlmostEqual(left, 50)
        self.assertAlmostEqual(top, 130)

    def test_relocated_images_and_utf8(self):
        self.write(image_object("old/location/asset.png"), encoding="utf-8-sig")
        report = convert(self.src, self.out, asset_dir=self.folder)
        self.assertEqual(len(report.slides[0]["objects"]), 1)

    def test_text_alpha(self):
        self.write(text_object("半透明").replace("透明度=0.0", "透明度=50.0"))
        convert(self.src, self.out)
        xml = Presentation(self.out).slides[0].shapes[0]._element.xml
        self.assertIn('a:alpha val="50000"', xml)

    def test_control_tags_rejected(self):
        self.write(text_object("<#ff0000>赤"))
        with self.assertRaises(ConversionError):
            convert(self.src, self.out)
        self.assertFalse(self.out.exists())


    def test_exo_rgb_color(self):
        self.write(text_object("青").replace("color=ffffff", "color=0000ff"))
        convert(self.src, self.out)
        run = Presentation(self.out).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        self.assertEqual(str(run.font.color.rgb), "0000FF")

    def test_missing_font_warns(self):
        self.write(text_object("字体").replace("font=MS Gothic", "font=Nonexistent_Font_123456"))
        with self.assertRaises(ConversionError):
            convert(self.src, self.out)
        self.assertIn("代用", self.out.with_suffix(".report.json").read_text("utf8"))


    def test_editable_rectangle_geometry_color_alpha(self):
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        self.write(figure_object())
        report = convert(self.src, self.out, frames=[1])
        self.assertFalse(report.warnings)
        p = Presentation(self.out)
        shape = p.slides[0].shapes[0]
        self.assertEqual(shape.shape_type, MSO_SHAPE_TYPE.AUTO_SHAPE)
        self.assertEqual(shape.auto_shape_type, MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        unit = p.slide_width / 1280
        self.assertAlmostEqual(shape.width/unit, 300, places=3)
        self.assertAlmostEqual(shape.height/unit, 150, places=3)
        self.assertAlmostEqual(shape.left/unit, 590, places=3)
        self.assertAlmostEqual(shape.top/unit, 185, places=3)
        self.assertEqual(shape.rotation, 30)
        self.assertEqual(str(shape.fill.fore_color.rgb), "B0B0FF")
        self.assertIn('a:alpha val="75000"', shape._element.xml)
        self.assertIsNone(shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style"))

    def test_background_canvas_and_layer_order(self):
        self.write(figure_object(kind=0, size=100, aspect=0)
                   .replace("拡大率=150", "拡大率=100").replace("X=100", "X=0")
                   .replace("Y=-100", "Y=0").replace("回転=30", "回転=0")
                   +text_object("背景の前の文字", ident=1, layer=2))
        convert(self.src, self.out, frames=[1])
        p = Presentation(self.out)
        bg, text = p.slides[0].shapes
        self.assertEqual((bg.left, bg.top), (0,0))
        self.assertEqual((bg.width, bg.height), (p.slide_width, p.slide_height))
        self.assertEqual(text.text, "背景の前の文字")

    def test_tall_rectangle(self):
        self.write(figure_object(aspect=50))
        report = convert(self.src, self.out, frames=[1])
        shape = report.slides[0]["objects"][0]
        self.assertEqual((shape["width_px"], shape["height_px"]), (150,300))

    def test_unsupported_figures_do_not_silently_become_rectangles(self):
        for index, body in enumerate([
            figure_object(kind=1), figure_object(line=5),
            figure_object().replace("name=", "name=custom"),
            figure_object().replace("サイズ=200", "サイズ=200,400,1"),
        ]):
            self.write(body)
            output = self.folder / f"unsupported-{index}.pptx"
            with self.assertRaises(ConversionError):
                convert(self.src, output)
            self.assertFalse(output.exists())

    def effect_image(self, effects, rotation=0):
        body = image_object().replace('回転=15', f'回転={rotation}')
        source, draw = body.split('[1.1]')
        return source + ''.join(f'[1.{i}]\n{effect}\n' for i, effect in enumerate(effects, 1)) + f'[1.{len(effects)+1}]' + draw

    def test_clip_retains_bitmap_and_original_rotation_origin(self):
        self.write(self.effect_image([
            '_name=クリッピング\n左=30\n右=10\n上=10\n下=20\n中心の位置を変更=0'
        ], rotation=90))
        report = convert(self.src, self.out, frames=[1])
        self.assertFalse(report.warnings)
        prs = Presentation(self.out)
        shape = prs.slides[0].shapes[0]
        unit = prs.slide_width/1280
        self.assertEqual(shape.image.size, (100,60))
        self.assertEqual((shape.crop_left,shape.crop_right), (.3,.1))
        self.assertAlmostEqual(shape.crop_top, 10/60, places=5)
        self.assertAlmostEqual(shape.crop_bottom, 20/60, places=5)
        for actual, expected in zip((shape.left,shape.top,shape.width,shape.height), (790,400,120,60)):
            self.assertAlmostEqual(actual/unit, expected, places=3)

    def test_clip_recenter(self):
        self.write(self.effect_image([
            '_name=クリッピング\n左=30\n右=10\n上=10\n下=20\n中心の位置を変更=1'
        ], rotation=90))
        report = convert(self.src, self.out, frames=[1])
        obj = report.slides[0]['objects'][0]
        self.assertEqual((obj['left_px'],obj['top_px']), (780,380))

    def test_flip_and_clip_order_with_asymmetric_pixels(self):
        im = Image.new('RGBA', (100,60), (255,0,0,255))
        im.paste((0,0,255,255), (50,0,100,60))
        im.save(self.folder/'asset.png')
        crop = '_name=クリッピング\n左=30\n右=10\n上=0\n下=0\n中心の位置を変更=0'
        flip = '_name=反転\n左右反転=1\n上下反転=1\n輝度反転=0\n色相反転=0\n透明度反転=0'
        for i, effects in enumerate(([crop,flip], [flip,crop], [flip,flip])):
            self.write(self.effect_image(effects))
            out = self.folder/f'order-{i}.pptx'
            report = convert(self.src, out, frames=[1])
            shape = Presentation(out).slides[0].shapes[0]
            self.assertFalse(report.warnings)
            self.assertEqual(shape.image.size, (100,60))
            with Image.open(BytesIO(shape.image.blob)) as full:
                self.assertEqual(full.getpixel((0,0))[:3], (255,0,0))
                self.assertEqual(full.getpixel((99,0))[:3], (0,0,255))
            if i < 2:
                self.assertEqual(shape._element.spPr.xfrm.get('flipH'), '1')
                self.assertEqual(shape._element.spPr.xfrm.get('flipV'), '1')
                self.assertEqual((shape.crop_left,shape.crop_right), ((.3,.1) if i==0 else (.1,.3)))
                self.assertEqual(report.slides[0]['objects'][0]['left_px'], 760 if i==0 else 800)
            else:
                self.assertIsNone(shape._element.spPr.xfrm.get('flipH'))
                self.assertIsNone(shape._element.spPr.xfrm.get('flipV'))

    def test_invalid_clips_and_unsupported_inversions_are_explicit(self):
        for i, effect in enumerate([
            '_name=クリッピング\n左=100', '_name=クリッピング\n上=-1',
            '_name=クリッピング\n左=0,20,1',
            '_name=反転\n輝度反転=1', '_name=反転\n左右反転=2',
        ]):
            self.write(self.effect_image([effect]))
            with self.assertRaises(ConversionError):
                convert(self.src, self.folder/f'invalid-effect-{i}.pptx', frames=[1])

    def test_text_crop_does_not_rasterize_or_silently_disappear(self):
        self.write(text_object('編集可能')+'[0.2]\n_name=クリッピング\n左=5\n')
        with self.assertRaises(ConversionError):
            convert(self.src, self.out, frames=[1])
        report = convert(self.src, self.folder/'text-approx.pptx', frames=[1], allow_approx=True)
        self.assertTrue(report.warnings)
        self.assertEqual(Presentation(self.folder/'text-approx.pptx').slides[0].shapes[0].text, '編集可能')

if __name__ == "__main__":
    unittest.main(verbosity=2)
