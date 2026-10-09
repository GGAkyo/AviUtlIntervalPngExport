"""Regression checks for snapshot ownership, handoff and asset resolution."""
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import gui
from exo_to_pptx import convert
from pptx import Presentation


class HandoffTests(unittest.TestCase):
    def test_output_skips_existing_deck_and_report(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "発表_editable.pptx"
            path.touch()
            path.with_name("発表_editable_2.report.json").touch()
            self.assertEqual(gui.available_output(path).name, "発表_editable_3.pptx")

    def test_command_line_preserves_unicode_and_spaces(self):
        options = gui.parse_options(["--source", "E:/作業 場所/scene.exo",
            "--output", "E:/中間 発表/発表.pptx", "--source-base", "E:/中間 発表",
            "--temporary-source"])
        self.assertEqual(options.output.name, "発表.pptx")
        self.assertTrue(options.temporary_source)

    def test_owned_snapshot_is_cleaned_without_deleting_other_files(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            folder = base / ".runtime" / "snapshots" / "epe123.tmp"
            folder.mkdir(parents=True)
            source = folder / "scene.exo"
            source.touch()
            marker = folder / "snapshot.owner"
            marker.write_bytes(b"EditablePptxExport snapshot v1\n")
            other = folder / "keep.txt"
            other.write_text("preserve")
            with patch.object(gui, "__file__", str(base / "gui.py")):
                gui.cleanup_snapshot(source)
            self.assertFalse(source.exists())
            self.assertFalse(marker.exists())
            self.assertEqual(other.read_text(), "preserve")

    def test_unowned_exo_is_never_deleted(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "scene.exo"
            source.write_text("original")
            with self.assertRaises(ValueError):
                gui.cleanup_snapshot(source)
            self.assertEqual(source.read_text(), "original")

    def test_snapshot_keeps_project_relative_asset_paths(self):
        example = Path(__file__).resolve().parent / "examples" / "sample.exo"
        with tempfile.TemporaryDirectory() as directory:
            snapshot = Path(directory) / "scene.exo"
            snapshot.write_bytes(example.read_bytes())
            result = Path(directory) / "snapshot.pptx"
            report = convert(snapshot, result, source_base=example.parent)
            self.assertFalse(report.warnings)
            deck = Presentation(result)
            self.assertEqual(len(deck.slides), 2)
            self.assertTrue(any(shape.shape_type == 13 for shape in deck.slides[0].shapes))


if __name__ == "__main__":
    unittest.main()
