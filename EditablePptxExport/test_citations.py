"""Citation numbering, persistence, export and IPC regression tests."""
from pathlib import Path
import struct
import tempfile
import unittest
from citations import resolve, pack, unpack, CitationError
from citation_gui import Row, Snapshot, make_updates, encode_updates, read_snapshot, import_entries
from exo_to_pptx import convert, ConversionError
from test_converter import HEADER, text_object
from pptx import Presentation

ENTRIES = {'smith2024': 'Smith et al., First paper, 2024.', 'tanaka2025': '田中ほか, 第二の論文, 2025.'}

class CitationTests(unittest.TestCase):
    def test_import_updates_all_matching_footnotes_and_preserves_missing_keys(self):
        import json
        entries = dict(ENTRIES)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'references.json'
            path.write_text(json.dumps({'smith2024': 'Updated bibliography'}), encoding='utf-8')
            self.assertEqual(import_entries(path, entries), 1)
            self.assertEqual(entries['tanaka2025'], ENTRIES['tanaka2025'])
            template = r'\ref{smith2024}'
            old = pack(template, '[1] ' + ENTRIES['smith2024'], ENTRIES)
            snapshot = Snapshot((1,2,3,4), 0, [Row(0,0,0,59,1,-1,old), Row(1,0,60,89,1,-1,old)])
            changes, rendered, _ = make_updates(snapshot, {0:template, 1:template}, entries)
            self.assertEqual(len(changes), 2)
            self.assertEqual(rendered, {0:'[1] Updated bibliography', 1:'[1] Updated bibliography'})
            path.write_text(json.dumps({'smith2024': 'Updated again'}), encoding='utf-8')
            import_entries(path, entries)
            self.assertEqual(entries['smith2024'], 'Updated again')
            path.write_text('invalid json', encoding='utf-8')
            before = dict(entries)
            with self.assertRaises(ValueError): import_entries(path, entries)
            self.assertEqual(entries, before)

    def test_bibliography_contains_visible_text_only(self):
        entries = {**ENTRIES, 'unused': 'Unused bibliography ' * 500}
        templates = {0:r'\cite{smith2024}', 1:r'\bibliography'}
        snapshot = Snapshot((1,2,3,4), 0, [Row(i,0,0,59,i,-1,t) for i,t in templates.items()])
        changes, _, _ = make_updates(snapshot, templates, entries)
        bibliography = next(text for index, _, text in changes if index == 1)
        self.assertEqual(bibliography, '[1] ' + ENTRIES['smith2024'])
        self.assertNotIn('<?', bibliography)

    def test_external_state_survives_reordering_and_export(self):
        from citation_state import load_state, update_state, save_state, recover
        templates = {0:r'A\cite{smith2024}', 1:r'B\cite{tanaka2025}', 2:r'\ref{smith2024}'}
        rows = [(0,0,1,templates[0]), (1,30,1,templates[1]), (2,0,2,templates[2])]
        rendered = resolve(rows, ENTRIES)[0]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'project.citations.json'
            state = update_state(load_state(path), 0, rows, templates, rendered, ENTRIES)
            save_state(path, state)
            reordered = [(10,30,1,rendered[0]), (11,0,1,rendered[1]), (12,0,2,rendered[2])]
            recovered, entries, warnings = recover(reordered, load_state(path))
            self.assertFalse(warnings)
            result = resolve([(i,s,l,recovered[i]) for i,s,l,t in reordered], entries)[0]
            self.assertEqual(result[10], 'A[2]')
            self.assertEqual(result[11], 'B[1]')
            source = Path(directory)/'scene.exo'
            output = Path(directory)/'result.pptx'
            source.write_text(HEADER + ''.join(text_object(t, ident=i, start=s+1, end=90, layer=l) for i,s,l,t in reordered), encoding='utf-8')
            convert(source, output, frames=[31], citation_state=path)
            strings = [shape.text for shape in Presentation(output).slides[0].shapes]
            self.assertIn('A[2]', strings)
            self.assertIn('[2] '+ENTRIES['smith2024'], strings)

    def test_long_bibliography_no_longer_pays_metadata_cost(self):
        template = r'\ref{smith2024}'
        entries = {'smith2024': '論文情報'*200}
        snapshot = Snapshot((1,2,3,4),0,[Row(0,0,0,59,1,-1,template)])
        changes, _, _ = make_updates(snapshot,{0:template},entries)
        self.assertEqual(changes[0][2], '[1] '+entries['smith2024'])
        with self.assertRaises(CitationError):
            make_updates(snapshot,{0:template},{'smith2024':'文'*1024})

    def test_old_marker_is_removed_and_ambiguous_state_is_rejected(self):
        from citation_state import recover, load_state, update_state
        template = r'\cite{smith2024}'
        old = pack(template, '[1]', ENTRIES)
        snapshot = Snapshot((1,2,3,4),0,[Row(0,0,0,59,1,-1,old)])
        self.assertEqual(make_updates(snapshot,{0:template},ENTRIES)[0], [(0,0,'[1]')])
        state = {'version':1, 'scenes':{'0':{'entries':ENTRIES,'rows':[
            {'template':template,'rendered':'[1]'}, {'template':r'\cite{tanaka2025}','rendered':'[1]'}]}}}
        with self.assertRaises(CitationError): recover([(0,0,0,'[1]')], state)
        templates, _, _ = recover([(0,0,0,old)], {'version':1,'scenes':{}})
        self.assertEqual(templates[0], template)

    def test_initial_occurrence_and_reordering(self):
        rows = [(0, 31, 1, r'A\cite{smith2024}'), (1, 1, 2, r'B\cite{tanaka2025}'),
                (2, 61, 3, r'脚注\ref{smith2024}'), (3, 1, 4, r'\bibliography')]
        rendered, numbers, _, _ = resolve(rows, ENTRIES)
        self.assertEqual(numbers, {'tanaka2025': 1, 'smith2024': 2})
        self.assertEqual(rendered[0], 'A[2]')
        self.assertIn('[2] Smith', rendered[2])
        self.assertEqual(rendered[3].splitlines()[0], '[1] ' + ENTRIES['tanaka2025'])
        rows[0] = (0, 0, 1, rows[0][3])
        self.assertEqual(resolve(rows, ENTRIES)[1], {'smith2024': 1, 'tanaka2025': 2})

    def test_duplicate_and_multiple_keys(self):
        result, numbers, _, _ = resolve([(0,1,1,r'\cite{smith2024,tanaka2025,smith2024} / \cite{smith2024}')],ENTRIES)
        self.assertEqual(result[0], '[1, 2] / [1]')
        self.assertEqual(len(numbers), 2)

    def test_refs_follow_citations_even_if_footnote_is_earlier(self):
        result, numbers, _, _ = resolve([(0,0,0,r'\ref{smith2024}'),(1,1,0,r'\cite{tanaka2025}')],ENTRIES)
        self.assertEqual(numbers, {'tanaka2025':1,'smith2024':2})

    def test_marker_round_trip_and_renumber_after_reopen(self):
        template = '関連研究' + r'\cite{smith2024}'
        managed = pack(template, '関連研究[2]', ENTRIES)
        self.assertEqual(unpack(managed)[0], template)
        result, numbers, _, _ = resolve([(0,0,0,managed)],None)
        self.assertEqual(result[0], '関連研究[1]')
        self.assertEqual(numbers, {'smith2024':1})

    def test_windows_multiline_text_survives_aviutl_and_exo_normalization(self):
        template = '段落\r\n' + r'\cite{smith2024}'
        managed = pack(template, '段落\r\n[1]', ENTRIES)
        self.assertEqual(unpack(managed)[0], '段落\n' + r'\cite{smith2024}')
        self.assertEqual(unpack(managed.replace('\n', '\r\n'))[2], '段落\n[1]')

    def test_direct_edits_are_detected_and_editor_can_recover_template(self):
        managed = pack(r'\cite{smith2024}', '[1]', ENTRIES) + '変更'
        with self.assertRaises(CitationError): unpack(managed)
        self.assertEqual(unpack(managed,verify=False)[0], r'\cite{smith2024}')

    def test_missing_and_malformed_keys_do_not_silently_render(self):
        for text in [r'\cite{unknown}', r'\cite{}', r'\cite{smith2024', r'\ref{bad key}']:
            with self.subTest(text=text), self.assertRaises(CitationError): resolve([(0,0,0,text)],ENTRIES)

    def test_reference_conflicts_are_explicit_and_catalog_can_override(self):
        a = pack(r'\cite{smith2024}','[1]',{'smith2024':'A'})
        b = pack(r'\cite{smith2024}','[1]',{'smith2024':'B'})
        with self.assertRaises(CitationError): resolve([(0,0,0,a),(1,1,0,b)])
        self.assertEqual(resolve([(0,0,0,a),(1,1,0,b)],{'smith2024':'C'})[2]['smith2024'],'C')

    def test_utf16_limit_and_lua_are_rejected(self):
        for template in ['漢'*1024+r'\cite{smith2024}', '<?return 1?>'+r'\cite{smith2024}', '😀'*600+r'\cite{smith2024}']:
            with self.assertRaises(CitationError): pack(template,template.replace(r'\cite{smith2024}','[1]'),ENTRIES)

    def test_restore_preserves_keys_without_catalog(self):
        template = r'\cite{smith2024}'
        text = pack(template,'[1]',ENTRIES)
        snapshot = Snapshot((1,2,3,4),0,[Row(0,0,0,59,1,-1,text)])
        changes, _, _ = make_updates(snapshot,{0:template},{},restore=True)
        self.assertEqual(changes, [(0,0,template)])

    def test_midpoint_updates_are_rejected(self):
        snapshot = Snapshot((1,2,3,4),0,[Row(0,0,0,59,1,0,r'\cite{smith2024}')])
        with self.assertRaises(CitationError): make_updates(snapshot,{0:r'\cite{smith2024}'},ENTRIES)

    def test_update_packet_keeps_unicode_and_session_token(self):
        snapshot = Snapshot((1,2,3,4),0,[])
        packet = encode_updates(snapshot,[(8,0,'日本語😀[1]')])
        self.assertEqual(struct.unpack_from('<6I',packet),(1,1,2,3,4,1))
        index, filter_index, length = struct.unpack_from('<3I',packet,24)
        self.assertEqual((index,filter_index),(8,0))
        self.assertEqual(length,len('日本語😀[1]'.encode('utf-16-le'))//2)
        self.assertEqual(packet[36:].decode('utf-16-le'),'日本語😀[1]')

    def test_snapshot_packet_validation(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'citation.input'
            text = '日本語'.encode('utf-16-le')
            data = struct.pack('<8I',0x45504349,1,1,2,3,4,0,1)+struct.pack('<6iI',8,0,0,59,1,-1,3)+text
            path.write_bytes(data)
            self.assertEqual(read_snapshot(path).rows[0].text,'日本語')
            path.write_bytes(data[:-1])
            with self.assertRaises(CitationError):read_snapshot(path)

    def test_pptx_resolves_plain_and_managed_text_and_renumbers_globally(self):
        import json
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory); src = base/'scene.exo'; output = base/'result.pptx'
            template = r'本文\cite{smith2024}'
            managed = pack(template,'本文[2]',ENTRIES)
            content = HEADER+text_object(managed,ident=0,start=1,end=30)
            content += text_object(r'\cite{tanaka2025}',ident=1,start=31,end=90)
            content += text_object(r'\ref{smith2024}',ident=2,start=1,end=90,layer=3)
            src.write_text(content,encoding='utf-8')
            (base/'references.json').write_text(json.dumps(ENTRIES,ensure_ascii=False),encoding='utf-8')
            report=convert(src,output,frames=[31])
            strings=[s.text for s in Presentation(output).slides[0].shapes]
            self.assertIn('[2]',strings) # numbering uses the whole project, not sampled slides
            self.assertIn('[1] '+ENTRIES['smith2024'],strings)
            self.assertFalse(report.warnings)

    def test_default_catalog_does_not_replace_managed_bibliography(self):
        managed = pack(r'\ref{smith2024}', '[1] New bibliography', {'smith2024':'New bibliography'})
        result, _, _, _ = resolve([(0,0,0,managed)],defaults={'smith2024':'Old bibliography'})
        self.assertEqual(result[0], '[1] New bibliography')
        result, _, _, _ = resolve([(0,0,0,managed)],entries={'smith2024':'Explicit bibliography'})
        self.assertEqual(result[0], '[1] Explicit bibliography')

    def test_bad_citation_blocks_approximate_export_as_well(self):
        with tempfile.TemporaryDirectory() as directory:
            src=Path(directory)/'bad.exo'; out=Path(directory)/'bad.pptx'
            src.write_text(HEADER+text_object(r'\cite{unknown}'),encoding='utf-8')
            with self.assertRaises(ConversionError):convert(src,out,allow_approx=True)
            self.assertFalse(out.exists())

if __name__ == '__main__':
    unittest.main()
